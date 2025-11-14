"""
主视图 / Main hub view for classroom selection
"""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import os
from datetime import datetime
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    import openpyxl.styles
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False


class MainView(ttk.Frame):
    """主视图框架 / Main view frame for displaying all classrooms"""

    def __init__(self, parent, controller, **kwargs):
        """
        初始化主视图
        / Initialize the main view
        
        Args:
            parent: 父级窗口 / Parent window
            controller: 应用控制器 / Application controller
        """
        super().__init__(parent, **kwargs)
        self.controller = controller
        self.configure(padding="20")

        # 标题 / Title
        title_label = ttk.Label(
            self, text="课堂管理系统 - 班级选择", font=("Arial", 18, "bold")
        )
        title_label.pack(pady=(0, 20))

        # 说明文本 / Description
        description_label = ttk.Label(self, text="请选择要管理的班级：")
        description_label.pack(fill=tk.X, pady=(0, 15))

        # 班级选择区域 / Classroom selection area
        self.selection_frame = ttk.Frame(self)
        self.selection_frame.pack(fill=tk.BOTH, expand=True, pady=10)
        self._max_cols = 3
        self._last_row_count = 0

        self.selected_class_var = tk.StringVar()
        self.class_buttons = {}
        self._populate_classroom_buttons()

        # 底部控制区域 / Bottom control area
        control_frame = ttk.Frame(self)
        control_frame.pack(fill=tk.X, pady=(20, 0))

        # 当前选择显示 / Current selection display
        self.selection_display = ttk.Label(
            control_frame, text="未选择班级", foreground="blue", font=("Arial", 12)
        )
        self.selection_display.pack(side=tk.LEFT, expand=True)

        # 按钮容器 / Button container
        button_frame = ttk.Frame(control_frame)
        button_frame.pack(side=tk.RIGHT)

        # 新增班级按钮 / Add new classroom button
        add_class_button = ttk.Button(
            button_frame, text="新增班级", command=self.add_classroom
        )
        add_class_button.pack(side=tk.LEFT, padx=5)

        # 退出登录 / Logout
        logout_button = ttk.Button(
            button_frame, text="退出登录", command=self.logout
        )
        logout_button.pack(side=tk.LEFT, padx=5)

        # 全屏切换 / Toggle fullscreen
        fullscreen_button = ttk.Button(
            button_frame, text="切换全屏", command=self.toggle_fullscreen
        )
        fullscreen_button.pack(side=tk.LEFT, padx=5)

    def _populate_classroom_buttons(self) -> None:
        """
        填充班级按钮 / Populate classroom buttons
        """
        # 清空现有网格
        for child in self.selection_frame.grid_slaves():
            child.destroy()

        # 重置之前的权重配置
        for index in range(self._last_row_count):
            self.selection_frame.rowconfigure(index, weight=0)
        for index in range(self._max_cols):
            self.selection_frame.columnconfigure(index, weight=0)

        self.class_buttons.clear()

        classrooms = self.controller.data_store.get_all_classrooms()
        if not classrooms:
            empty_label = ttk.Label(self.selection_frame, text="暂无班级，请先创建。")
            empty_label.grid(row=0, column=0, padx=10, pady=10)
            self._last_row_count = 1
            return

        button_row = 0
        button_col = 0

        for classroom in classrooms:
            # 创建按钮 / Create button for each classroom
            button = ttk.Button(
                self.selection_frame,
                text=f"{classroom.name}\n({len(classroom.students)}名学生)",
                command=lambda class_id=classroom.class_id: self.on_class_selected(
                    class_id
                ),
                width=20,
            )
            button.grid(row=button_row, column=button_col, padx=5, pady=5, sticky="nsew")

            # 绑定右键菜单 / Bind right-click menu
            button.bind(
                "<Button-3>",
                lambda e, cid=classroom.class_id: self.show_context_menu(e, cid),
            )

            self.class_buttons[classroom.class_id] = button

            button_col += 1
            if button_col >= self._max_cols:
                button_col = 0
                button_row += 1

        # 配置网格权重 / Configure grid weights
        row_count = button_row + 1
        for i in range(row_count):
            self.selection_frame.rowconfigure(i, weight=1)
        for j in range(self._max_cols):
            self.selection_frame.columnconfigure(j, weight=1)

        self._last_row_count = row_count

    def show_context_menu(self, event, class_id: str) -> None:
        """
        显示班级右键菜单 / Show classroom context menu
        """
        context_menu = tk.Menu(self, tearoff=0)
        context_menu.add_command(label="修改班级名字", command=lambda: self.edit_classroom_by_id(class_id))
        context_menu.add_command(label="删除班级", command=lambda: self.delete_classroom_by_id(class_id))
        context_menu.add_separator()
        context_menu.add_command(label="导出Excel", command=lambda: self.export_class_excel(class_id))
        
        try:
            context_menu.tk_popup(event.x_root, event.y_root)
        finally:
            context_menu.grab_release()

    def edit_classroom_by_id(self, class_id: str) -> None:
        """
        根据ID修改班级名字 / Edit classroom name by ID
        """
        self.selected_class_var.set(class_id)
        classroom = self.controller.data_store.get_classroom(class_id)
        self.selection_display.config(
            text=f"已选择: {classroom.name} ({len(classroom.students)}名学生)"
        )
        self.edit_classroom()

    def delete_classroom_by_id(self, class_id: str) -> None:
        """
        根据ID删除班级 / Delete classroom by ID
        """
        self.selected_class_var.set(class_id)
        classroom = self.controller.data_store.get_classroom(class_id)
        self.selection_display.config(
            text=f"已选择: {classroom.name} ({len(classroom.students)}名学生)"
        )
        self.delete_classroom()

    def on_class_selected(self, class_id: str) -> None:
        """
        处理班级选择 - 直接进入班级详情页
        / Handle classroom selection - Directly enter class detail page
        
        Args:
            class_id: 班级ID / Classroom ID
        """
        self.selected_class_var.set(class_id)
        classroom = self.controller.data_store.get_classroom(class_id)
        self.selection_display.config(
            text=f"已选择: {classroom.name} ({len(classroom.students)}名学生)"
        )
        # 直接进入班级详情页 / Directly enter class detail page
        self.controller.set_current_classroom(class_id)
        self.controller.show_frame("ClassDetailView")

    def enter_class_detail(self) -> None:
        """
        进入班级详情视图
        / Enter the class detail view
        """
        class_id = self.selected_class_var.get()
        if not class_id:
            messagebox.showwarning("未选择班级", "请先选择一个班级。")
            return

        self.controller.set_current_classroom(class_id)
        self.controller.show_frame("ClassDetailView")

    def toggle_fullscreen(self) -> None:
        """
        切换全屏模式 / Toggle fullscreen mode
        """
        current_state = self.controller.state()
        if current_state == 'zoomed':
            self.controller.state('normal')
        else:
            self.controller.state('zoomed')

    def logout(self) -> None:
        """
        退出登录，返回登录视图
        / Logout and return to login view
        """
        self.selected_class_var.set("")
        self.selection_display.config(text="未选择班级")
        self.controller.show_frame("LoginView")

    def reset(self) -> None:
        """重置主视图 / Reset main view"""
        self.selected_class_var.set("")
        self.selection_display.config(text="未选择班级")
        self.refresh_classrooms()

    def refresh_classrooms(self) -> None:
        """
        刷新班级按钮显示 / Refresh classroom buttons display
        """
        self._populate_classroom_buttons()

    def add_classroom(self) -> None:
        """
        添加新班级 / Add a new classroom
        """
        dialog = tk.Toplevel(self)
        dialog.title("添加班级")
        dialog.geometry("300x150")
        dialog.transient(self.winfo_toplevel())
        dialog.grab_set()

        ttk.Label(dialog, text="班级名字：").pack(pady=(10, 0), padx=10)
        name_entry = ttk.Entry(dialog, width=30)
        name_entry.pack(pady=5, padx=10, fill=tk.X)
        name_entry.focus()

        def confirm():
            name = name_entry.get().strip()
            if not name:
                messagebox.showwarning("输入错误", "班级名字不能为空。")
                return
            
            self.controller.data_store.add_classroom(name)
            dialog.destroy()
            self.refresh_classrooms()
            messagebox.showinfo("成功", f"班级 {name} 已添加。")

        button_frame = ttk.Frame(dialog)
        button_frame.pack(pady=10)

        ttk.Button(button_frame, text="确定", command=confirm).pack(side=tk.LEFT, padx=5)
        ttk.Button(button_frame, text="取消", command=dialog.destroy).pack(side=tk.LEFT, padx=5)

    def edit_classroom(self) -> None:
        """
        修改班级名字 / Edit classroom name
        """
        class_id = self.selected_class_var.get()
        if not class_id:
            messagebox.showwarning("未选择班级", "请先选择一个班级。")
            return

        classroom = self.controller.data_store.get_classroom(class_id)
        
        dialog = tk.Toplevel(self)
        dialog.title("修改班级")
        dialog.geometry("300x150")
        dialog.transient(self.winfo_toplevel())
        dialog.grab_set()

        ttk.Label(dialog, text="班级名字：").pack(pady=(10, 0), padx=10)
        name_entry = ttk.Entry(dialog, width=30)
        name_entry.insert(0, classroom.name)
        name_entry.pack(pady=5, padx=10, fill=tk.X)
        name_entry.focus()

        def confirm():
            new_name = name_entry.get().strip()
            if not new_name:
                messagebox.showwarning("输入错误", "班级名字不能为空。")
                return
            
            self.controller.data_store.update_classroom_name(class_id, new_name)
            dialog.destroy()
            self.refresh_classrooms()
            messagebox.showinfo("成功", f"班级已改名为 {new_name}。")

        button_frame = ttk.Frame(dialog)
        button_frame.pack(pady=10)

        ttk.Button(button_frame, text="确定", command=confirm).pack(side=tk.LEFT, padx=5)
        ttk.Button(button_frame, text="取消", command=dialog.destroy).pack(side=tk.LEFT, padx=5)

    def delete_classroom(self) -> None:
        """
        删除班级 / Delete a classroom
        """
        class_id = self.selected_class_var.get()
        if not class_id:
            messagebox.showwarning("未选择班级", "请先选择一个班级。")
            return

        classroom = self.controller.data_store.get_classroom(class_id)
        
        if messagebox.askyesno(
            "确认删除",
            f"确定要删除班级 {classroom.name} 吗？\n此操作不可撤销。",
        ):
            self.controller.data_store.remove_classroom(class_id)
            self.selected_class_var.set("")
            self.selection_display.config(text="未选择班级")
            self.refresh_classrooms()
            messagebox.showinfo("成功", f"班级 {classroom.name} 已删除。")

    def export_class_excel(self, class_id: str) -> None:
        """
        导出班级小码币数据到Excel文件
        / Export class coin data to Excel file
        """
        if not EXCEL_AVAILABLE:
            messagebox.showerror("功能不可用", "Excel导出功能需要安装openpyxl库。")
            return
            
        classroom = self.controller.data_store.get_classroom(class_id)
        if not classroom:
            messagebox.showerror("错误", "班级数据获取失败。")
            return
            
        # 选择保存位置 / Choose save location
        filename = f"{classroom.name}_小码币数据_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        file_path = filedialog.asksaveasfilename(
            title="保存Excel文件",
            initialfile=filename,
            defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")]
        )
        
        if not file_path:
            return  # 用户取消了保存
            
        try:
            # 创建工作簿 / Create workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "小码币统计"
            
            # 设置标题样式 / Set title style
            title_font = Font(name='微软雅黑', size=14, bold=True)
            header_font = Font(name='微软雅黑', size=12, bold=True)
            header_fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")
            header_alignment = Alignment(horizontal='center', vertical='center')
            
            # 合并标题单元格 / Merge title cells
            ws.merge_cells('A1:B1')
            ws['A1'] = f"{classroom.name} - 小码币数据统计"
            ws['A1'].font = title_font
            ws['A1'].alignment = header_alignment
            
            # 设置表头 / Set headers
            ws['A2'] = "学生名字"
            ws['B2'] = "总小码币数量"
            
            # 应用表头样式 / Apply header styles
            for cell in ['A2', 'B2']:
                ws[cell].font = header_font
                ws[cell].fill = header_fill
                ws[cell].alignment = header_alignment
            
            # 填充数据 / Fill data
            row = 3
            data_store = self.controller.data_store
            for student in classroom.get_all_students():
                total_coins = data_store.get_student_total(student)
                ws[f'A{row}'] = student.name
                ws[f'B{row}'] = total_coins
                row += 1
            
            # 设置列宽 / Set column widths
            ws.column_dimensions['A'].width = 20
            ws.column_dimensions['B'].width = 15
            
            # 添加边框 / Add borders
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            for row in range(1, row):
                for col in ['A', 'B']:
                    ws[f'{col}{row}'].border = thin_border
            
            # 保存文件 / Save file
            wb.save(file_path)
            messagebox.showinfo("导出成功", f"Excel文件已保存到：\n{file_path}")
            
        except Exception as e:
            messagebox.showerror("导出失败", f"导出Excel文件时发生错误：\n{str(e)}")
