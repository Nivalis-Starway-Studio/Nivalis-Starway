"""
Excel数据存储模块 / Excel data storage module
"""

import os
from typing import Dict, List
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from .models import Classroom, Student


class ExcelDataStore:
    """Excel数据存储和管理 / Excel data store and management"""
    
    def __init__(self, file_path: str = "classroom_data.xlsx"):
        """
        初始化Excel数据存储
        / Initialize Excel data store
        
        Args:
            file_path: Excel文件路径 / Excel file path
        """
        self.file_path = file_path
        # 确保文件路径是绝对路径
        if not os.path.isabs(self.file_path):
            # 使用脚本所在目录
            script_dir = os.path.dirname(os.path.abspath(__file__))
            project_dir = os.path.dirname(script_dir)
            self.file_path = os.path.join(project_dir, self.file_path)
    
    def save_classrooms(self, classrooms: Dict[str, Classroom]) -> bool:
        """
        保存所有班级数据到Excel
        / Save all classroom data to Excel
        
        Args:
            classrooms: 班级字典 / Dictionary of classrooms
            
        Returns:
            是否成功 / Success status
        """
        try:
            wb = Workbook()
            # 删除默认的Sheet
            if "Sheet" in wb.sheetnames:
                wb.remove(wb["Sheet"])
            
            for class_id, classroom in classrooms.items():
                # 为每个班级创建一个工作表
                ws = wb.create_sheet(title=classroom.name)
                
                # 设置表头样式
                header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                header_font = Font(bold=True, color="FFFFFF", size=11)
                
                # 写入表头
                headers = ["学生名字", "累计币"]
                
                # 添加周数列（最多显示52周）
                for week in range(1, 53):
                    headers.append(f"第{week}周")
                
                for col_idx, header in enumerate(headers, start=1):
                    cell = ws.cell(row=1, column=col_idx, value=header)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                
                # 写入学生数据
                for row_idx, student in enumerate(classroom.get_all_students(), start=2):
                    # 学生名字
                    ws.cell(row=row_idx, column=1, value=student.name)
                    # 累计币
                    ws.cell(row=row_idx, column=2, value=student.cumulative_coins)
                    
                    # 周币数据
                    for week_num, coins in student.week_coins.items():
                        col_idx = 2 + week_num  # 第1周在第3列
                        ws.cell(row=row_idx, column=col_idx, value=coins)
                
                # 调整列宽
                ws.column_dimensions["A"].width = 15
                ws.column_dimensions["B"].width = 12
                for col_num in range(3, 55):
                    # 正确转换列号为列字母
                    col_letter = self._get_column_letter(col_num)
                    ws.column_dimensions[col_letter].width = 10
            
            wb.save(self.file_path)
            return True
        except Exception as e:
            print(f"保存Excel文件失败 / Failed to save Excel: {e}")
            return False
    
    def load_classrooms(self) -> Dict[str, Classroom]:
        """
        从Excel加载所有班级数据
        / Load all classroom data from Excel
        
        Returns:
            班级字典 / Dictionary of classrooms
        """
        classrooms = {}
        
        if not os.path.exists(self.file_path):
            return classrooms
        
        try:
            wb = load_workbook(self.file_path)
            
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                
                # 从表头读取周数列
                # 假设格式: 学生名字 | 累计币 | 第1周 | 第2周 | ...
                headers = [cell.value for cell in ws[1]]
                
                # 提取班级ID（简单处理，可能需要改进）
                # 这里假设班级名称格式如"一年级A班"
                classroom_name = sheet_name
                class_id = self._extract_class_id(classroom_name)
                
                classroom = Classroom(name=classroom_name, class_id=class_id)
                
                # 读取学生数据（从第2行开始）
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row[0]:  # 如果名字为空，跳过
                        continue
                    
                    student_name = row[0]
                    cumulative_coins = row[1] if row[1] is not None else 0
                    
                    # 创建学生对象
                    student = Student(
                        name=student_name,
                        cumulative_coins=cumulative_coins,
                        week_coins={}
                    )
                    
                    # 读取周币数据
                    for col_idx in range(2, len(headers)):
                        header = headers[col_idx]
                        if header and header.startswith("第") and header.endswith("周"):
                            # 提取周数
                            week_str = header[1:-1]
                            try:
                                week_num = int(week_str)
                                coins = row[col_idx] if col_idx < len(row) and row[col_idx] is not None else 0
                                if coins > 0:  # 只保存非零的周币
                                    student.week_coins[week_num] = coins
                            except ValueError:
                                continue
                    
                    classroom.add_student(student)
                
                classrooms[class_id] = classroom
            
            wb.close()
        except Exception as e:
            print(f"加载Excel文件失败 / Failed to load Excel: {e}")
        
        return classrooms
    
    def _get_column_letter(self, col_num: int) -> str:
        """
        将列号转换为列字母（1->A, 2->B, ..., 27->AA）
        / Convert column number to letter (1->A, 2->B, ..., 27->AA)
        
        Args:
            col_num: 列号 / Column number
            
        Returns:
            列字母 / Column letter
        """
        result = ""
        while col_num > 0:
            col_num, remainder = divmod(col_num - 1, 26)
            result = chr(65 + remainder) + result
        return result

    def _extract_class_id(self, classroom_name: str) -> str:
        """
        从班级名称提取班级ID
        / Extract class ID from classroom name
        
        Args:
            classroom_name: 班级名称 / Classroom name
            
        Returns:
            班级ID / Class ID
        """
        # 简单的映射规则
        mapping = {
            "一年级A班": "1A",
            "一年级B班": "1B",
            "二年级A班": "2A",
            "二年级B班": "2B",
            "三年级A班": "3A",
            "三年级B班": "3B",
            "四年级A班": "4A",
        }
        return mapping.get(classroom_name, classroom_name)
